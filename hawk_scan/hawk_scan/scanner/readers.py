import os

MAX_EXTRACTED_BYTES = 10 * 1024 * 1024  # 10 MB extracted text cap
MAX_XLSX_ROWS = 50_000
MAX_PDF_PAGES = 200


def read_text(file_path: str) -> str:
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")
    with open(file_path, "rb") as f:
        return f.read(MAX_EXTRACTED_BYTES).decode("utf-8", errors="replace")


def read_pdf(file_path: str) -> str:
    import warnings
    import logging
    import threading
    import PyPDF2
    logging.getLogger("PyPDF2").setLevel(logging.ERROR)
    parts: list[str] = []
    total = 0
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for i, page in enumerate(reader.pages):
                if i >= MAX_PDF_PAGES or total >= MAX_EXTRACTED_BYTES:
                    break
                result: list[str | None] = [None]
                def _extract(p=page):
                    try:
                        result[0] = p.extract_text()
                    except Exception:
                        pass
                t = threading.Thread(target=_extract)
                t.daemon = True
                t.start()
                t.join(timeout=3)
                if t.is_alive():
                    continue
                if result[0]:
                    parts.append(result[0])
                    total += len(result[0])
    return "\n".join(parts)


def read_docx(file_path: str) -> str:
    from docx import Document
    doc = Document(file_path)
    return "\n".join(p.text for p in doc.paragraphs)


def read_xlsx(file_path: str) -> str:
    import warnings
    from openpyxl import load_workbook
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(file_path, data_only=True, read_only=True)
        parts = []
        total = 0
        try:
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                row_count = 0
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            val = str(cell.value)
                            parts.append(val)
                            total += len(val)
                    row_count += 1
                    if row_count >= MAX_XLSX_ROWS or total >= MAX_EXTRACTED_BYTES:
                        break
                if total >= MAX_EXTRACTED_BYTES:
                    break
        finally:
            wb.close()
    return "\n".join(parts)


def read_pptx(file_path: str) -> str:
    from pptx import Presentation
    prs = Presentation(file_path)
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        parts.append(text)
    return "\n".join(parts)


def read_image_ocr(file_path: str) -> str:
    try:
        import warnings
        import numpy as np
        import cv2
        import pytesseract
        from PIL import Image, ImageEnhance
    except (ImportError, AttributeError):
        return ""

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        image = Image.open(file_path)
        if image.mode in ("P", "PA"):
            image = image.convert("RGBA")
        if image.mode == "RGBA":
            bg = Image.new("RGB", image.size, (255, 255, 255))
            bg.paste(image, mask=image.split()[3])
            image = bg
        grayscale = image.convert("L")
        enhancer = ImageEnhance.Contrast(grayscale)
        contrasted = enhancer.enhance(2.0)
        thresholded = contrasted.point(lambda x: 0 if x < 100 else 255)
        arr = np.array(thresholded)
        denoised = cv2.fastNlMeansDenoising(arr, None, h=10,
                                             templateWindowSize=7,
                                             searchWindowSize=21)
        enhanced = Image.fromarray(denoised)
        return pytesseract.image_to_string(enhanced)


EXTENSION_MAP = {
    ".txt": read_text,
    ".csv": read_text,
    ".pdf": read_pdf,
    ".docx": read_docx,
    ".xlsx": read_xlsx,
    ".pptx": read_pptx,
    ".png": read_image_ocr,
    ".jpg": read_image_ocr,
    ".jpeg": read_image_ocr,
    ".gif": read_image_ocr,
    ".bmp": read_image_ocr,
}

SCANNABLE_EXTENSIONS = set(EXTENSION_MAP.keys())


def read_file(file_path: str) -> str:
    _, ext = os.path.splitext(file_path)
    ext = ext.lower()
    reader = EXTENSION_MAP.get(ext, read_text)
    return reader(file_path)
