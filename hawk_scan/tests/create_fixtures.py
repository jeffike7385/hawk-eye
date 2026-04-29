import os
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from PIL import Image, ImageDraw
from reportlab.pdfgen import canvas as rl_canvas

FIXTURES = os.path.join(os.path.dirname(__file__), "fixtures")
os.makedirs(FIXTURES, exist_ok=True)

# Plain text
with open(os.path.join(FIXTURES, "sample.txt"), "w") as f:
    f.write("Employee SSN: 123-45-6789\nEmail: test@example.com\n")

# Word doc
doc = Document()
doc.add_paragraph("Employee SSN: 234-56-7890")
doc.add_paragraph("Contact: hr@company.com")
doc.save(os.path.join(FIXTURES, "sample.docx"))

# Excel
wb = Workbook()
ws = wb.active
ws["A1"] = "Name"
ws["B1"] = "SSN"
ws["A2"] = "Jane Doe"
ws["B2"] = "345-67-8901"
wb.save(os.path.join(FIXTURES, "sample.xlsx"))

# PowerPoint
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Confidential Report"
slide.placeholders[1].text = "SSN: 456-78-9012\nEmail: user@corp.com"
prs.save(os.path.join(FIXTURES, "sample.pptx"))

# Image with text
img = Image.new("RGB", (400, 100), color="white")
draw = ImageDraw.Draw(img)
draw.text((10, 40), "SSN 567-89-0123", fill="black")
img.save(os.path.join(FIXTURES, "sample.png"))

# PDF
pdf_path = os.path.join(FIXTURES, "sample.pdf")
c = rl_canvas.Canvas(pdf_path)
c.drawString(72, 720, "Employee SSN: 678-90-1234")
c.drawString(72, 700, "Email: finance@company.com")
c.save()

print(f"Fixtures created in {FIXTURES}")
